"""UI flow tests with Streamlit's AppTest (no browser, no keys, no internet): the real app.py runs with fake engines
(tests/ui/harness_app.py). Run: pytest -q"""
import os
import sys

import pytest
from streamlit.testing.v1 import AppTest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(os.path.dirname(HERE)))
import core  # noqa: E402

LINKS = "\n".join(["https://shop-a.example/", "https://shop-a.example/about/", "https://dir-one.example/listing/shop-a-dental",
                   "https://dir-two.example/profile/shop-a", "https://dir-three.example/shop-a", "https://dir-four.example/biz/shop-a",
                   "dir-five.example/x", "not a link", "https://shop-a.example/"])


def md(at):
    return " ".join(m.value for m in at.markdown)


@pytest.fixture
def at():
    a = AppTest.from_file(os.path.join(HERE, "harness_app.py"), default_timeout=60)
    a.run()
    assert not a.exception, a.exception
    return a


def test_first_screen_is_honest(at):
    page = md(at)
    assert "No method tested yet" in page            # a typed key is never shown as "working"
    assert "Not tested yet" in page and "Off" in page
    assert "Daily allowance" in page and "Sign-up credit" in page
    check = [b for b in at.button if b.label.startswith("▶")][0]
    assert check.label == "▶  Check 0 links" and check.disabled


def test_links_check_cards_filter_and_status(at):
    at.text_area[0].set_value(LINKS).run()
    page = md(at)
    assert "7 links ready" in page and "1 duplicate removed" in page and "1 line not a link" in page
    check = [b for b in at.button if b.label.startswith("▶")][0]
    assert check.label == "▶  Check 7 links" and not check.disabled
    check.click().run()
    assert not at.exception, at.exception
    warn = " ".join(w.value for w in at.warning)
    assert "1 link could not be checked" in warn and "7 of 7 links checked" in warn
    labels = {b.key: b.label for b in at.button if b.key and b.key.startswith("cardbtn_")}
    assert labels["cardbtn_0"].startswith("**3**") and labels["cardbtn_2"].startswith("**2**") and labels["cardbtn_4"].startswith("**1**")
    assert "Working" in md(at)                        # Serper answered in the check -> now proven to work
    at.button(key="cardbtn_2").click().run()          # card = filter
    assert any("Showing 2 of 7 links" in c.value and "Not found on Google" in c.value for c in at.caption)
    at.button(key="cardbtn_2").click().run()          # again = all
    assert any("Showing 7 of 7 links" in c.value for c in at.caption)
    assert any(b.label.startswith("🔁  Re-check the 1 link") for b in at.button)


def test_stopped_check_keeps_what_it_found(at):
    r = core.CheckResult(url="https://shop-a.example/", verdict=core.V.INDEXED, detail="x", engine="Serper", checked_at="t")
    at.session_state["partial"] = {"total": 5, "results": {r.url: r}, "at": "2026-10-21 09:00 UTC"}
    at.run()
    warn = " ".join(w.value for w in at.warning)
    assert "stopped after 1 of 5 links" in warn and "Not complete" in warn


def test_exports_are_safe():
    import pandas as pd
    df = pd.DataFrame({"Link": ["https://a.example/x,y", "=HYPERLINK(\"http://evil\")"], "Result": [core.V.INDEXED, core.V.NOTFOUND]})
    csv = core.to_csv(df).decode("utf-8-sig")
    assert '"https://a.example/x,y"' in csv and "'=HYPERLINK" in csv
    tsv = core.to_tsv(df)
    assert "\t" in tsv and "'=HYPERLINK" in tsv
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(core.to_excel(df, {"Links checked": 2})))
    ws = wb["Index Check"]
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row) and ws.auto_filter.ref


def test_stopped_recheck_keeps_every_earlier_row(at):
    rs = [core.CheckResult(url=f"https://s{i}.example/", verdict=core.V.UNKNOWN if i == 1 else core.V.INDEXED, detail="x",
                           engine="Serper", checked_at="t") for i in range(3)]
    at.session_state["results"] = rs
    at.session_state["run_meta"] = {"at": "t", "total": 3, "used": {}, "dead": {}}
    fixed = core.CheckResult(url="https://s1.example/", verdict=core.V.INDEXED, detail="y", engine="Serper", checked_at="t2")
    at.session_state["partial"] = {"total": 1, "results": {fixed.url: fixed}, "at": "t2", "merge": True}
    at.run()
    res = at.session_state["results"]
    assert len(res) == 3 and res[1].verdict == core.V.INDEXED


def test_password_in_any_language():
    a = AppTest.from_file(os.path.join(HERE, "harness_app.py"), default_timeout=60)
    a.secrets["APP_PASSWORD"] = "pässwörd-پاس"
    a.run()
    a.text_input[0].set_value("wrong-ü").run()
    assert not a.exception and any("Wrong password" in e.value for e in a.error)
    a.text_input[0].set_value("pässwörd-پاس").run()
    assert not a.exception and a.session_state["authed"]
